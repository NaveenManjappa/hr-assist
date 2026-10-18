from collections.abc import Iterable
from dataclasses import dataclass
from typing import TYPE_CHECKING, Protocol

from hr_assist.corpus import ManifestEntry

if TYPE_CHECKING:
    from pptx.shapes.base import BaseShape
    from pptx.slide import Slide
    from pptx.table import Table


@dataclass(frozen=True)
class Page:
    """One unit of a Source Document as Markdown: a PDF page, or a deck slide numbered from 1."""

    number: int
    text: str


class Parser(Protocol):
    def parse(self, entry: ManifestEntry) -> list[Page]: ...


class PdfParser:
    """Loads each PDF page as Markdown, which keeps headings and pipe tables (ADR 0001)."""

    def parse(self, entry: ManifestEntry) -> list[Page]:
        from langchain_pymupdf4llm import PyMuPDF4LLMLoader

        documents = PyMuPDF4LLMLoader(entry.path, mode="page").load()
        return [Page(number=doc.metadata["page"] + 1, text=doc.page_content) for doc in documents]


class DeckParser:
    """Renders each slide as Markdown: `## title`, body text, tables, then the speaker notes.

    The notes hold policy content that isn't on the slide, so they are kept (ADR 0001).
    Slide number placeholders, pictures and charts carry no text worth embedding.
    """

    def parse(self, entry: ManifestEntry) -> list[Page]:
        from pptx import Presentation

        slides = Presentation(str(entry.path)).slides
        return [Page(number, _render_slide(slide)) for number, slide in enumerate(slides, start=1)]


def _render_slide(slide: "Slide") -> str:
    from pptx.enum.shapes import PP_PLACEHOLDER
    from pptx.shapes.graphfrm import GraphicFrame

    title = ""
    blocks: list[str] = []
    text_lines: list[str] = []
    for shape in _shapes(slide.shapes):
        placeholder = shape.placeholder_format.type if shape.is_placeholder else None
        if placeholder == PP_PLACEHOLDER.SLIDE_NUMBER:
            continue
        if placeholder in (PP_PLACEHOLDER.TITLE, PP_PLACEHOLDER.CENTER_TITLE) and not title:
            title = " ".join(_paragraphs(shape))
            continue
        if isinstance(shape, GraphicFrame) and shape.has_table:
            if text_lines:
                blocks.append("\n".join(text_lines))
                text_lines = []
            blocks.append(_markdown_table(shape.table))
        else:
            text_lines.extend(_paragraphs(shape))
    if text_lines:
        blocks.append("\n".join(text_lines))
    notes_frame = slide.notes_slide.notes_text_frame if slide.has_notes_slide else None
    if notes_frame is not None and (notes := notes_frame.text.strip()):
        blocks.append(f"Speaker notes: {notes}")
    return "\n\n".join([f"## {title}", *blocks])


def _shapes(shapes: "Iterable[BaseShape]") -> "list[BaseShape]":
    """Flattens group shapes, keeping the slide's reading order."""
    from pptx.shapes.group import GroupShape

    flat: list[BaseShape] = []
    for shape in shapes:
        if isinstance(shape, GroupShape):
            flat.extend(_shapes(shape.shapes))
        else:
            flat.append(shape)
    return flat


def _paragraphs(shape: "BaseShape") -> list[str]:
    from pptx.shapes.autoshape import Shape

    if not isinstance(shape, Shape):
        return []
    return [p.text.strip() for p in shape.text_frame.paragraphs if p.text.strip()]


def _markdown_table(table: "Table") -> str:
    rows = [[_cell(cell.text) for cell in row.cells] for row in table.rows]
    lines = [f"|{'|'.join(row)}|" for row in rows]
    lines.insert(1, f"|{'|'.join('---' for _ in rows[0])}|")
    return "\n".join(lines)


def _cell(text: str) -> str:
    return " ".join(text.split()).replace("|", "\\|")
